"""Quarentena de lançamentos: valida os Excel em dados/entrada_quarentena/ antes de entrarem em dados/banco_de_dados/.

Uso:
    python validar_entrada.py                 # valida, gera pendências e promove ficheiros sem bloqueantes
    python validar_entrada.py --sem-promover  # só valida e gera pendências (não move nada)

Saídas em saidas/relatorios_auditoria/: pendencias_lancamentos.xlsx, pendencias_lancamentos.html, log_quarentena.txt.
A correção é sempre feita no Sankhya, com nova exportação; o Excel exportado não é editado à mão.
"""
import argparse
import datetime
import html
import json
import os
import shutil
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from fontes_dados import PASTA_DADOS, PASTA_QUARENTENA, PASTA_RELATORIOS, listar_ficheiros_quarentena
from regras_lancamento import AVISO, BLOQUEANTE, configurar_stdout_utf8, validar_ficheiro

CAMINHO_XLSX = os.path.join(PASTA_RELATORIOS, "pendencias_lancamentos.xlsx")
CAMINHO_HTML = os.path.join(PASTA_RELATORIOS, "pendencias_lancamentos.html")
CAMINHO_LOG = os.path.join(PASTA_RELATORIOS, "log_quarentena.txt")
PASTA_SUBSTITUIDOS = os.path.join(PASTA_QUARENTENA, "substituidos")

STATUS_OPCOES = ["Pendente", "Corrigido no Sankhya", "Justificado"]

COLUNAS = [
    ("Severidade", "severidade", 13),
    ("Regra", "regra", 28),
    ("Ficheiro", "ficheiro", 30),
    ("Linha Excel", "linha", 8),
    ("Nro Único", "nr_unico", 12),
    ("Data", "data", 11),
    ("Parceiro", "parceiro", 30),
    ("Operador", "operador", 20),
    ("Valor (R$)", "valor", 14),
    ("Placa ERP", "placa_erp", 26),
    ("Histórico", "historico", 45),
    ("Detalhe", "detalhe", 28),
    ("Ação esperada", "acao", 55),
    ("Status", "status", 20),
    ("Observação do operador", "observacao", 40),
]

PREENCHIMENTO = {
    BLOQUEANTE: PatternFill("solid", fgColor="FDE2E2"),
    AVISO: PatternFill("solid", fgColor="FFF4D6"),
}


def chave_pendencia(erro: dict) -> tuple:
    if erro["nr_unico"] != "N/D":
        return (erro["nr_unico"], erro["regra"])
    return (erro["ficheiro"], str(erro["linha"]), erro["regra"])


def ler_status_anteriores() -> dict:
    """Status e Observação já preenchidos pelo operador na planilha anterior."""
    if not os.path.isfile(CAMINHO_XLSX):
        return {}
    try:
        ws = load_workbook(CAMINHO_XLSX, read_only=True)["Pendências"]
    except Exception as exc:
        print(f"[AVISO] Não foi possível ler a planilha anterior ({exc}); Status/Observação recomeçam.")
        return {}
    linhas = ws.iter_rows(values_only=True)
    cabecalho = [str(c) if c is not None else "" for c in next(linhas, [])]
    idx = {nome: i for i, nome in enumerate(cabecalho)}
    necessarias = ["Nro Único", "Regra", "Ficheiro", "Linha Excel", "Status", "Observação do operador"]
    if any(n not in idx for n in necessarias):
        return {}
    anteriores = {}
    for linha in linhas:
        def campo(nome):
            v = linha[idx[nome]]
            return "" if v is None else str(v)
        erro = {"nr_unico": campo("Nro Único") or "N/D", "regra": campo("Regra"),
                "ficheiro": campo("Ficheiro"), "linha": campo("Linha Excel")}
        anteriores[chave_pendencia(erro)] = (campo("Status"), campo("Observação do operador"))
    return anteriores


def gravar_xlsx(erros: list, resumo_ficheiros: list, gerado_em: str) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Pendências"
    ws.append([c[0] for c in COLUNAS])
    for celula in ws[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="1B4D3E")
        celula.alignment = Alignment(vertical="center", wrap_text=True)

    col_valor = [c[1] for c in COLUNAS].index("valor") + 1
    for erro in erros:
        ws.append([erro.get(campo, "") if campo != "valor" else erro.get("valor") for _, campo, _ in COLUNAS])
        linha = ws.max_row
        for celula in ws[linha]:
            celula.fill = PREENCHIMENTO[erro["severidade"]]
            celula.alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row=linha, column=col_valor).number_format = '#,##0.00;[Red]-#,##0.00'

    for i, (_, _, largura) in enumerate(COLUNAS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura
    ws.freeze_panes = "A2"
    ultima_col = get_column_letter(len(COLUNAS))
    ws.auto_filter.ref = f"A1:{ultima_col}{max(ws.max_row, 2)}"

    col_status = get_column_letter([c[1] for c in COLUNAS].index("status") + 1)
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUS_OPCOES) + '"', allow_blank=False)
    dv.error = "Escolha um status da lista."
    dv.prompt = "Pendente / Corrigido no Sankhya / Justificado"
    ws.add_data_validation(dv)
    dv.add(f"{col_status}2:{col_status}{max(ws.max_row, 2) + 200}")

    wr = wb.create_sheet("Resumo")
    wr.append(["Quarentena de lançamentos - gerado em", gerado_em])
    wr.append([])
    wr.append(["Ficheiro", "Bloqueantes", "Avisos", "Decisão"])
    for celula in wr[3]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="1B4D3E")
    for r in resumo_ficheiros:
        wr.append([r["ficheiro"], r["bloqueantes"], r["avisos"], r["decisao"]])
    wr.append([])
    wr.append(["Bloqueante: impede a entrada em dados/banco_de_dados até corrigir no Sankhya e reexportar."])
    wr.append(["Aviso: não impede; confirmar e marcar Status (Corrigido no Sankhya / Justificado)."])
    for col, largura in zip("ABCD", (45, 12, 10, 60)):
        wr.column_dimensions[col].width = largura

    os.makedirs(PASTA_RELATORIOS, exist_ok=True)
    try:
        wb.save(CAMINHO_XLSX)
        return CAMINHO_XLSX
    except PermissionError:
        alternativo = CAMINHO_XLSX.replace(".xlsx", f"_{datetime.datetime.now():%Y%m%d_%H%M%S}.xlsx")
        wb.save(alternativo)
        print(f"[AVISO] {os.path.basename(CAMINHO_XLSX)} está aberto no Excel; gravado em {os.path.basename(alternativo)}.")
        return alternativo


def promover(caminho: str) -> str:
    """Move o ficheiro para dados/banco_de_dados; um homónimo existente vai para substituidos/."""
    nome = os.path.basename(caminho)
    destino = os.path.join(PASTA_DADOS, nome)
    detalhe = ""
    if os.path.exists(destino):
        os.makedirs(PASTA_SUBSTITUIDOS, exist_ok=True)
        antigo = os.path.join(PASTA_SUBSTITUIDOS, f"{datetime.datetime.now():%Y%m%d_%H%M%S}_{nome}")
        shutil.move(destino, antigo)
        detalhe = f" (versão anterior movida para substituidos/{os.path.basename(antigo)})"
    shutil.move(caminho, destino)
    return f"PROMOVIDO para dados/banco_de_dados{detalhe}"


def registar_log(linhas: list) -> None:
    os.makedirs(PASTA_RELATORIOS, exist_ok=True)
    with open(CAMINHO_LOG, "a", encoding="utf-8") as f:
        f.write("\n".join(linhas) + "\n\n")


def gerar_html(erros: list, resumo_ficheiros: list, gerado_em: str, sem_promover: bool) -> None:
    bloqueantes = [e for e in erros if e["severidade"] == BLOQUEANTE]
    linhas_bloqueadas = {(e["ficheiro"], e["linha"]): abs(e["valor"] or 0) for e in bloqueantes}
    valor_risco = sum(linhas_bloqueadas.values())
    promovidos = sum(1 for r in resumo_ficheiros if r["decisao"].startswith("PROMOVIDO"))
    retidos = sum(1 for r in resumo_ficheiros if r["decisao"].startswith("RETIDO"))
    modo = "Simulação (--sem-promover): nada foi movido" if sem_promover else "Execução real"

    dados = json.dumps(
        [{k: e.get(k) for k in ("severidade", "regra", "ficheiro", "linha", "nr_unico", "data", "parceiro",
                                 "operador", "valor", "placa_erp", "historico", "detalhe", "acao", "status")}
         for e in erros],
        ensure_ascii=False,
    ).replace("</", "<\\/")
    ficheiros_json = json.dumps(resumo_ficheiros, ensure_ascii=False).replace("</", "<\\/")
    valor_risco_txt = f"R$ {valor_risco:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

    pagina = f"""<!DOCTYPE html>
<html lang="pt-BR">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Agrovia - Quarentena de Lançamentos</title>
<link rel="stylesheet" href="https://cdnjs.cloudflare.com/ajax/libs/font-awesome/6.4.0/css/all.min.css">
<style>
    :root {{
        --bg-color: #0f1e17;
        --card-bg: #162c22;
        --border-color: rgba(46, 204, 113, 0.25);
        --text-primary: #f8fafc;
        --text-secondary: #94a3b8;
        --accent-blue: #38bdf8;
        --accent-purple: #a3e635;
        --accent-pink: #2ecc71;
        --accent-green: #34d399;
        --accent-yellow: #fbbf24;
        --accent-orange: #fb923c;
        --accent-red: #f87171;
    }}
    body {{ font-family: 'Inter', -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif; background: var(--bg-color); color: var(--text-primary); margin: 0; padding: 15px; }}
    .container {{ max-width: 1400px; margin: 0 auto; display: flex; flex-direction: column; gap: 20px; }}
    .card {{ background: var(--card-bg); border: 1px solid var(--border-color); border-radius: 12px; padding: 15px; box-shadow: 0 4px 12px rgba(0,0,0,0.3); box-sizing: border-box; }}
    header {{ display: flex; justify-content: space-between; align-items: center; flex-wrap: wrap; gap: 15px; }}
    header h1 {{ margin: 0; font-size: 20px; background: linear-gradient(90deg, var(--accent-blue), var(--accent-purple)); -webkit-background-clip: text; background-clip: text; -webkit-text-fill-color: transparent; }}
    header p {{ margin: 5px 0 0; font-size: 12px; color: var(--text-secondary); }}
    .author-signature {{ font-size: 11px; color: var(--accent-purple); margin-top: 4px; font-weight: 600; }}
    .actions {{ display: flex; gap: 10px; flex-wrap: wrap; }}
    .btn {{ background: rgba(56,189,248,0.1); border: 1px solid rgba(56,189,248,0.3); color: var(--accent-blue); padding: 8px 14px; border-radius: 8px; cursor: pointer; font-weight: 600; font-size: 12px; }}
    .btn.active, .btn:hover {{ background: var(--accent-blue); color: #0f1e17; }}
    .btn-print {{ background: rgba(163,230,53,0.15); border-color: rgba(163,230,53,0.4); color: var(--accent-purple); }}
    .btn-print:hover {{ background: var(--accent-purple); color: #0f1e17; }}
    .kpis {{ display: grid; grid-template-columns: repeat(5, 1fr); gap: 15px; }}
    .kpi {{ position: relative; overflow: hidden; }}
    .kpi::after {{ content: ''; position: absolute; top: 0; left: 0; width: 4px; height: 100%; background: var(--k); }}
    .kpi-title {{ font-size: 11px; color: var(--text-secondary); font-weight: 600; text-transform: uppercase; }}
    .kpi-value {{ font-size: 22px; font-weight: 700; margin-top: 8px; }}
    .kpi-sub {{ font-size: 11px; color: var(--text-secondary); margin-top: 4px; }}
    .grid2 {{ display: grid; grid-template-columns: 1fr 1fr; gap: 15px; }}
    .card-title {{ font-size: 14px; font-weight: 600; margin-bottom: 12px; color: var(--accent-purple); }}
    table {{ width: 100%; border-collapse: collapse; font-size: 12px; }}
    th, td {{ padding: 8px 10px; border-bottom: 1px solid var(--border-color); text-align: left; vertical-align: top; }}
    th {{ background: rgb(19,44,34); color: var(--text-secondary); font-size: 11px; text-transform: uppercase; letter-spacing: 0.4px; }}
    .table-wrap th {{ position: sticky; top: 0; z-index: 1; }}
    .nowrap {{ white-space: nowrap; }}
    tr:hover td {{ background: rgba(46,204,113,0.05); }}
    .num {{ text-align: right; white-space: nowrap; }}
    .badge {{ display: inline-block; padding: 3px 8px; border-radius: 6px; font-size: 10px; font-weight: 700; letter-spacing: 0.3px; white-space: nowrap; }}
    .badge-BLOQUEANTE {{ background: rgba(248,113,113,0.18); color: var(--accent-red); border: 1px solid rgba(248,113,113,0.5); }}
    .badge-AVISO {{ background: rgba(251,191,36,0.15); color: var(--accent-yellow); border: 1px solid rgba(251,191,36,0.45); }}
    .badge-ok {{ background: rgba(52,211,153,0.15); color: var(--accent-green); border: 1px solid rgba(52,211,153,0.45); }}
    .filters {{ display: flex; flex-wrap: wrap; gap: 10px; margin-bottom: 12px; align-items: center; }}
    .filters select, .filters input {{ background: rgba(10,23,17,0.9); color: var(--text-primary); border: 1px solid var(--border-color); border-radius: 6px; padding: 6px 10px; font-size: 12px; }}
    .filters input {{ flex: 1; min-width: 180px; }}
    .count {{ font-size: 12px; color: var(--text-secondary); }}
    .table-wrap {{ max-height: 620px; overflow: auto; }}
    .acao {{ color: var(--text-secondary); font-size: 11px; }}
    #tabelaPendencias .check, body.mobile-mode #tabelaPendencias td.check {{ display: none; }}
    .empty {{ text-align: center; color: var(--text-secondary); padding: 20px; }}
    footer {{ display: flex; justify-content: space-between; align-items: center; flex-wrap: wrap; gap: 15px; font-size: 12px; color: var(--text-secondary); }}
    .footer-left {{ display: flex; align-items: center; gap: 8px; color: var(--text-primary); font-weight: 600; }}
    .footer-left span {{ color: var(--accent-purple); font-weight: 700; }}
    .footer-right {{ font-size: 11px; background: rgba(10,23,17,0.5); padding: 6px 12px; border-radius: 6px; border: 1px solid var(--border-color); }}

    @media (max-width: 1100px) {{ .kpis {{ grid-template-columns: repeat(3, 1fr); }} }}
    @media (max-width: 760px) {{
        .kpis {{ grid-template-columns: repeat(2, 1fr); }}
        .grid2 {{ grid-template-columns: 1fr; }}
    }}
    body.mobile-mode .container {{ max-width: 430px; }}
    body.mobile-mode .kpis {{ grid-template-columns: repeat(2, 1fr); }}
    body.mobile-mode .grid2 {{ grid-template-columns: 1fr; }}
    @media (max-width: 760px) {{
        #tabelaPendencias thead {{ display: none; }}
        #tabelaPendencias tr {{ display: block; border: 1px solid var(--border-color); border-radius: 10px; margin-bottom: 10px; padding: 6px; }}
        #tabelaPendencias td {{ display: flex; justify-content: space-between; gap: 10px; border: none; padding: 4px 6px; text-align: right; }}
        #tabelaPendencias td::before {{ content: attr(data-label); color: var(--text-secondary); font-size: 11px; text-align: left; font-weight: 600; }}
    }}
    body.mobile-mode #tabelaPendencias thead {{ display: none; }}
    body.mobile-mode #tabelaPendencias tr {{ display: block; border: 1px solid var(--border-color); border-radius: 10px; margin-bottom: 10px; padding: 6px; }}
    body.mobile-mode #tabelaPendencias td {{ display: flex; justify-content: space-between; gap: 10px; border: none; padding: 4px 6px; text-align: right; }}
    body.mobile-mode #tabelaPendencias td::before {{ content: attr(data-label); color: var(--text-secondary); font-size: 11px; text-align: left; font-weight: 600; }}

    @media print {{
        body {{ background: #fff; color: #000; padding: 0; font-size: 10px; }}
        .card {{ background: #fff; border: 1px solid #999; box-shadow: none; break-inside: avoid; }}
        header h1 {{ -webkit-text-fill-color: #000; color: #000; background: none; }}
        .actions, .filters, .no-print {{ display: none !important; }}
        .table-wrap {{ max-height: none; overflow: visible; }}
        th {{ background: #e5e5e5; color: #000; position: static; }}
        th, td {{ border-bottom: 1px solid #bbb; color: #000; }}
        .acao, .kpi-sub, .kpi-title, .author-signature, footer, .footer-left, .footer-left span {{ color: #000; }}
        .badge {{ border: 1px solid #000; color: #000; background: none; }}
        body:not(.mobile-mode) #tabelaPendencias .check {{ display: table-cell; }}
        tr {{ break-inside: avoid; }}
    }}
</style>
</head>
<body id="bodyTag">
<div class="container">
    <header class="card">
        <div>
            <h1><i class="fa-solid fa-shield-halved"></i> Agrovia - Quarentena de Lançamentos</h1>
            <p>Gerado em: {html.escape(gerado_em)} | {html.escape(modo)} | Correção sempre no Sankhya, com nova exportação</p>
            <div class="author-signature">Autor do Relatório: Samuel Oliveira Silva</div>
        </div>
        <div class="actions">
            <button class="btn active" onclick="setViewMode('pc', this)"><i class="fa-solid fa-desktop"></i> PC</button>
            <button class="btn" onclick="setViewMode('mobile', this)"><i class="fa-solid fa-mobile-screen"></i> Mobile</button>
            <button class="btn btn-print" onclick="window.print()"><i class="fa-solid fa-print"></i> Imprimir ordem de serviço</button>
        </div>
    </header>

    <section class="kpis">
        <div class="card kpi" style="--k: var(--accent-blue)"><div class="kpi-title">Total de erros</div><div class="kpi-value">{len(erros)}</div><div class="kpi-sub">bloqueantes + avisos</div></div>
        <div class="card kpi" style="--k: var(--accent-red)"><div class="kpi-title">Bloqueantes</div><div class="kpi-value">{len(bloqueantes)}</div><div class="kpi-sub">impedem a entrada no banco</div></div>
        <div class="card kpi" style="--k: var(--accent-yellow)"><div class="kpi-title">Avisos</div><div class="kpi-value">{len(erros) - len(bloqueantes)}</div><div class="kpi-sub">confirmar e justificar</div></div>
        <div class="card kpi" style="--k: var(--accent-orange)"><div class="kpi-title">Valor em risco</div><div class="kpi-value">{valor_risco_txt}</div><div class="kpi-sub">{len(linhas_bloqueadas)} lançamentos com bloqueante</div></div>
        <div class="card kpi" style="--k: var(--accent-purple)"><div class="kpi-title">Ficheiros</div><div class="kpi-value">{promovidos} / {retidos}</div><div class="kpi-sub">promovidos / retidos</div></div>
    </section>

    <section class="card">
        <div class="card-title"><i class="fa-solid fa-folder-open"></i> Decisão por ficheiro</div>
        <div class="table-wrap"><table id="tabelaFicheiros"><thead><tr><th>Ficheiro</th><th class="num">Bloqueantes</th><th class="num">Avisos</th><th>Decisão</th></tr></thead><tbody></tbody></table></div>
    </section>

    <section class="grid2">
        <div class="card">
            <div class="card-title"><i class="fa-solid fa-user-pen"></i> Por operador</div>
            <table id="tabelaOperador"><thead><tr><th>Operador</th><th class="num">Bloqueantes</th><th class="num">Avisos</th></tr></thead><tbody></tbody></table>
        </div>
        <div class="card">
            <div class="card-title"><i class="fa-solid fa-list-check"></i> Por regra</div>
            <table id="tabelaRegra"><thead><tr><th>Regra</th><th>Severidade</th><th class="num">Qtd.</th></tr></thead><tbody></tbody></table>
        </div>
    </section>

    <section class="card">
        <div class="card-title"><i class="fa-solid fa-triangle-exclamation"></i> Pendências a corrigir</div>
        <div class="filters">
            <select id="fSeveridade" aria-label="Filtrar severidade"><option value="">Todas as severidades</option><option>BLOQUEANTE</option><option>AVISO</option></select>
            <select id="fOperador" aria-label="Filtrar operador"><option value="">Todos os operadores</option></select>
            <select id="fFicheiro" aria-label="Filtrar ficheiro"><option value="">Todos os ficheiros</option></select>
            <input id="fBusca" type="search" placeholder="Procurar Nro Único, parceiro, histórico..." aria-label="Procurar">
            <span class="count" id="contagem"></span>
        </div>
        <div class="table-wrap">
            <table id="tabelaPendencias">
                <thead><tr>
                    <th class="check">OK</th><th>Severidade</th><th>Regra</th><th>Nro Único</th><th>Linha</th><th>Data</th>
                    <th>Parceiro</th><th>Operador</th><th class="num">Valor (R$)</th><th>Placa ERP / Detalhe</th><th>Ação esperada</th><th>Status</th>
                </tr></thead>
                <tbody></tbody>
            </table>
        </div>
    </section>

    <footer class="card">
        <div class="footer-left"><i class="fa-solid fa-shield-cat" style="color: var(--accent-purple);"></i> Gestão e Auditoria Técnica: <span>Samuel Oliveira Silva</span></div>
        <div class="footer-right"><i class="fa-solid fa-code-branch" style="color: var(--accent-blue);"></i> Agrovia Quarentena de Lançamentos • Gerado em: {html.escape(gerado_em)}</div>
    </footer>
</div>

<script>
const ERROS = {dados};
const FICHEIROS = {ficheiros_json};

function setViewMode(mode, btn) {{
    document.querySelectorAll('.actions .btn:not(.btn-print)').forEach(b => b.classList.remove('active'));
    btn.classList.add('active');
    document.getElementById('bodyTag').classList.toggle('mobile-mode', mode === 'mobile');
}}

function moeda(v) {{
    if (v === null || v === undefined) return 'sem valor';
    return v.toLocaleString('pt-BR', {{ minimumFractionDigits: 2, maximumFractionDigits: 2 }});
}}

function celula(tr, texto, label, classe) {{
    const td = document.createElement('td');
    td.textContent = texto;
    if (label) td.dataset.label = label;
    if (classe) td.className = classe;
    tr.appendChild(td);
    return td;
}}

function badge(td, sev) {{
    td.textContent = '';
    const s = document.createElement('span');
    s.className = 'badge badge-' + sev;
    s.textContent = sev;
    td.appendChild(s);
}}

function preencherOpcoes(id, valores) {{
    const sel = document.getElementById(id);
    [...new Set(valores)].sort().forEach(v => {{ const o = document.createElement('option'); o.textContent = v; sel.appendChild(o); }});
}}

function renderResumos() {{
    const tbF = document.querySelector('#tabelaFicheiros tbody');
    FICHEIROS.forEach(f => {{
        const tr = document.createElement('tr');
        celula(tr, f.ficheiro); celula(tr, f.bloqueantes, null, 'num'); celula(tr, f.avisos, null, 'num');
        const td = celula(tr, '');
        const s = document.createElement('span');
        s.className = 'badge ' + (f.decisao.startsWith('PROMOVIDO') || f.decisao.startsWith('APTO') ? 'badge-ok' : 'badge-BLOQUEANTE');
        s.textContent = f.decisao;
        td.appendChild(s);
        tbF.appendChild(tr);
    }});
    if (!FICHEIROS.length) tbF.innerHTML = '<tr><td colspan="4" class="empty">Nenhum ficheiro em dados/entrada_quarentena.</td></tr>';

    const porOp = {{}}, porRegra = {{}};
    ERROS.forEach(e => {{
        porOp[e.operador] = porOp[e.operador] || {{ BLOQUEANTE: 0, AVISO: 0 }};
        porOp[e.operador][e.severidade]++;
        const k = e.regra + '|' + e.severidade;
        porRegra[k] = (porRegra[k] || 0) + 1;
    }});
    const tbO = document.querySelector('#tabelaOperador tbody');
    Object.entries(porOp).sort((a, b) => (b[1].BLOQUEANTE - a[1].BLOQUEANTE) || (b[1].AVISO - a[1].AVISO)).forEach(([op, c]) => {{
        const tr = document.createElement('tr');
        celula(tr, op); celula(tr, c.BLOQUEANTE, null, 'num'); celula(tr, c.AVISO, null, 'num');
        tbO.appendChild(tr);
    }});
    const tbR = document.querySelector('#tabelaRegra tbody');
    Object.entries(porRegra).sort((a, b) => b[1] - a[1]).forEach(([k, n]) => {{
        const [regra, sev] = k.split('|');
        const tr = document.createElement('tr');
        celula(tr, regra); badge(celula(tr, ''), sev); celula(tr, n, null, 'num');
        tbR.appendChild(tr);
    }});
    if (!ERROS.length) {{
        tbO.innerHTML = '<tr><td colspan="3" class="empty">Sem pendências.</td></tr>';
        tbR.innerHTML = '<tr><td colspan="3" class="empty">Sem pendências.</td></tr>';
    }}
}}

function renderPendencias() {{
    const sev = document.getElementById('fSeveridade').value;
    const op = document.getElementById('fOperador').value;
    const fic = document.getElementById('fFicheiro').value;
    const busca = document.getElementById('fBusca').value.trim().toLowerCase();
    const tb = document.querySelector('#tabelaPendencias tbody');
    tb.innerHTML = '';
    const filtrados = ERROS.filter(e =>
        (!sev || e.severidade === sev) && (!op || e.operador === op) && (!fic || e.ficheiro === fic) &&
        (!busca || [e.nr_unico, e.parceiro, e.historico, e.regra, e.placa_erp, e.detalhe].join(' ').toLowerCase().includes(busca))
    );
    filtrados.forEach(e => {{
        const tr = document.createElement('tr');
        celula(tr, '[  ]', 'OK', 'check');
        badge(celula(tr, '', 'Severidade'), e.severidade);
        celula(tr, e.regra, 'Regra');
        celula(tr, e.nr_unico, 'Nro Único');
        celula(tr, e.ficheiro + ' : ' + e.linha, 'Ficheiro : linha');
        celula(tr, e.data, 'Data', 'nowrap');
        celula(tr, e.parceiro, 'Parceiro');
        celula(tr, e.operador, 'Operador');
        celula(tr, moeda(e.valor), 'Valor (R$)', 'num');
        celula(tr, e.placa_erp + (e.detalhe ? ' | ' + e.detalhe : ''), 'Placa / Detalhe');
        celula(tr, e.acao, 'Ação esperada', 'acao');
        celula(tr, e.status || 'Pendente', 'Status');
        tr.title = e.historico || '';
        tb.appendChild(tr);
    }});
    if (!filtrados.length) tb.innerHTML = '<tr><td colspan="12" class="empty">Nenhuma pendência com estes filtros.</td></tr>';
    document.getElementById('contagem').textContent = filtrados.length + ' de ' + ERROS.length + ' pendências';
}}

preencherOpcoes('fOperador', ERROS.map(e => e.operador));
preencherOpcoes('fFicheiro', ERROS.map(e => e.ficheiro));
['fSeveridade', 'fOperador', 'fFicheiro', 'fBusca'].forEach(id => {{
    const el = document.getElementById(id);
    el.addEventListener('input', renderPendencias);
    el.addEventListener('change', renderPendencias);
}});
renderResumos();
renderPendencias();
</script>
</body>
</html>
"""
    os.makedirs(PASTA_RELATORIOS, exist_ok=True)
    with open(CAMINHO_HTML, "w", encoding="utf-8") as f:
        f.write(pagina)


def main() -> int:
    configurar_stdout_utf8()
    parser = argparse.ArgumentParser(description="Valida os Excel em dados/entrada_quarentena antes de entrarem em dados/banco_de_dados.")
    parser.add_argument("--sem-promover", action="store_true", help="só valida e gera pendências; não move ficheiros")
    args = parser.parse_args()

    os.makedirs(PASTA_QUARENTENA, exist_ok=True)
    agora = datetime.datetime.now()
    gerado_em = f"{agora:%d/%m/%Y às %H:%M}"

    print("=" * 100)
    print(" AGROVIA - QUARENTENA DE LANÇAMENTOS" + (" (SIMULAÇÃO: --sem-promover)" if args.sem_promover else ""))
    print("=" * 100)

    ficheiros = listar_ficheiros_quarentena()
    if not ficheiros:
        print(" Nenhum ficheiro Excel em dados/entrada_quarentena/. Nada a validar.")
        return 0

    anteriores = ler_status_anteriores()
    erros_todos = []
    resumo_ficheiros = []
    log = [f"[{agora:%Y-%m-%d %H:%M:%S}] validar_entrada.py" + (" --sem-promover" if args.sem_promover else "")]

    for caminho in ficheiros:
        nome = os.path.basename(caminho)
        try:
            erros = validar_ficheiro(caminho)
        except Exception as exc:
            print(f" [X] {nome}: não foi possível ler ({exc}).")
            resumo_ficheiros.append({"ficheiro": nome, "bloqueantes": 1, "avisos": 0, "decisao": f"RETIDO: ficheiro ilegível ({exc})"})
            log.append(f"  - {nome} | RETIDO | ficheiro ilegível: {exc}")
            continue

        for erro in erros:
            status, observacao = anteriores.get(chave_pendencia(erro), ("", ""))
            erro["status"] = status or "Pendente"
            erro["observacao"] = observacao
        erros_todos.extend(erros)

        n_bloq = sum(1 for e in erros if e["severidade"] == BLOQUEANTE)
        n_aviso = len(erros) - n_bloq
        if n_bloq:
            decisao = f"RETIDO: {n_bloq} bloqueante(s) - corrigir no Sankhya e reexportar"
        elif args.sem_promover:
            decisao = "APTO (simulação: não movido)"
        else:
            try:
                decisao = promover(caminho)
            except OSError as exc:
                decisao = f"RETIDO: falha ao mover ({exc})"
        resumo_ficheiros.append({"ficheiro": nome, "bloqueantes": n_bloq, "avisos": n_aviso, "decisao": decisao})
        log.append(f"  - {nome} | bloqueantes={n_bloq} avisos={n_aviso} | {decisao}")
        print(f" {nome}: {n_bloq} bloqueante(s), {n_aviso} aviso(s) -> {decisao}")

    caminho_xlsx = gravar_xlsx(erros_todos, resumo_ficheiros, gerado_em)
    gerar_html(erros_todos, resumo_ficheiros, gerado_em, args.sem_promover)
    registar_log(log)

    print("-" * 100)
    print(f" Planilha do operador : saidas\\relatorios_auditoria\\{os.path.basename(caminho_xlsx)}")
    print(f" Resumo visual (HTML) : saidas\\relatorios_auditoria\\{os.path.basename(CAMINHO_HTML)}")
    print(f" Log                  : saidas\\relatorios_auditoria\\{os.path.basename(CAMINHO_LOG)}")
    print("=" * 100)

    return 2 if any(r["decisao"].startswith("RETIDO") for r in resumo_ficheiros) else 0


if __name__ == "__main__":
    sys.exit(main())
